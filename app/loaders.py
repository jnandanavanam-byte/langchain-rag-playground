"""
Document Loaders

Supported Formats:
- PDF
- DOCX
- XLSX
- CSV
- TXT

Converts all documents into LangChain Document objects.
"""

import os

import pandas as pd

from pypdf import PdfReader
from docx import Document as DocxDocument

from langchain_core.documents import Document


# ============================================================
# PDF Loader
# ============================================================

def load_pdf(file_path):
    """
    Load PDF document and create one Document per page.
    """

    documents = []

    reader = PdfReader(file_path)

    for page_num, page in enumerate(
        reader.pages,
        start=1
    ):

        text = page.extract_text()

        if text and text.strip():

            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": os.path.basename(file_path),
                        "page": page_num,
                        "file_type": "pdf"
                    }
                )
            )

    return documents


# ============================================================
# DOCX Loader
# ============================================================

def load_docx(file_path):
    """
    Load Word document.
    """

    doc = DocxDocument(file_path)

    text = "\n".join(
        para.text
        for para in doc.paragraphs
        if para.text.strip()
    )

    return [
        Document(
            page_content=text,
            metadata={
                "source": os.path.basename(file_path),
                "file_type": "docx"
            }
        )
    ]


# ============================================================
# TXT Loader
# ============================================================

def load_txt(file_path):
    """
    Load plain text file.
    """

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        text = file.read()

    return [
        Document(
            page_content=text,
            metadata={
                "source": os.path.basename(file_path),
                "file_type": "txt"
            }
        )
    ]


# ============================================================
# CSV Loader
# ============================================================

def load_csv(file_path):
    """
    Load CSV file.
    """

    df = pd.read_csv(file_path)

    content = df.to_string(index=False)

    return [
        Document(
            page_content=content,
            metadata={
                "source": os.path.basename(file_path),
                "rows": len(df),
                "file_type": "csv"
            }
        )
    ]


# ============================================================
# XLSX Loader
# ============================================================

def load_xlsx(file_path):
    """
    Load Excel workbook.
    Creates one document per worksheet.
    """

    documents = []

    excel_file = pd.ExcelFile(
        file_path,
        engine="openpyxl"
    )

    for sheet_name in excel_file.sheet_names:

        df = pd.read_excel(
            file_path,
            sheet_name=sheet_name,
            engine="openpyxl"
        )

        content = df.to_string(
            index=False
        )

        documents.append(
            Document(
                page_content=content,
                metadata={
                    "source": os.path.basename(file_path),
                    "sheet": sheet_name,
                    "rows": len(df),
                    "file_type": "xlsx"
                }
            )
        )

    return documents


# ============================================================
# Generic Loader
# ============================================================

def load_file(file_path):
    """
    Automatically detect and load file type.
    """

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":
        return load_pdf(file_path)

    if extension == ".docx":
        return load_docx(file_path)

    if extension == ".csv":
        return load_csv(file_path)

    if extension == ".xlsx":
        return load_xlsx(file_path)

    if extension == ".txt":
        return load_txt(file_path)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )


# ============================================================
# Bulk Loader
# ============================================================

def load_documents(file_paths):
    """
    Load all uploaded files.

    Returns:
        documents
        statistics
    """

    all_documents = []

    for file_path in file_paths:

        all_documents.extend(
            load_file(file_path)
        )

    stats = get_document_statistics(
        all_documents,
        len(file_paths)
    )

    return all_documents, stats


# ============================================================
# Statistics
# ============================================================

def get_document_statistics(
    documents,
    file_count
):
    """
    Generate metrics for Streamlit status panel.
    """

    total_characters = sum(
        len(doc.page_content)
        for doc in documents
    )

    pdf_pages = sum(
        1
        for doc in documents
        if doc.metadata.get(
            "file_type"
        ) == "pdf"
    )

    return {
        "files": file_count,
        "documents": len(documents),
        "pages": pdf_pages,
        "characters": total_characters
    }